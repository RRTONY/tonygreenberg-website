"""Rebuilds Migration-Status-Report.xlsx from NEXTJS-MIGRATION-TODO.md.

The spreadsheet is the team's at-a-glance copy of the main project file, so it
is generated, never edited by hand: run this after updating the todo file.
Three tabs, same look as before: Summary (per-phase counts as formulas over
the Full Detail tab), Full Detail (every checkbox), Remaining Work (open items
only).

Status rules: [x] is Done; an open item is Deferred when its text says it is
deferred, on hold or cancelled (or it sits in the cancelled Phase 10),
otherwise Remaining.

Usage: python3 scripts/build-migration-status-report.py
Needs openpyxl (pip install openpyxl). Open the file in Excel/Numbers to see
the formulas calculated, or run a LibreOffice recalc first.
"""

import datetime
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parent.parent
TODO = ROOT / "NEXTJS-MIGRATION-TODO.md"
OUT = ROOT / "Migration-Status-Report.xlsx"

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
PHASE_FILL = PatternFill("solid", fgColor="D1D5DB")
STATUS_FILL = {
    "Done": PatternFill("solid", fgColor="D9F2D9"),
    "Remaining": PatternFill("solid", fgColor="FDE2E2"),
    "Deferred": PatternFill("solid", fgColor="FEF3C7"),
}
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DEFERRED = re.compile(r"\b(defer(red)?|on hold|cancelled|canceled)\b", re.I)


def clean(md: str) -> str:
    """Markdown to plain text: links, bold/italic/code marks, strikethrough."""
    t = re.sub(r"~~(.*?)~~", r"\1 (done)", md)
    t = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", t)
    t = re.sub(r"(\*\*|__|`)", "", t)
    t = re.sub(r"(?<!\w)\*(?!\s)(.+?)(?<!\s)\*(?!\w)", r"\1", t)
    return re.sub(r"\s+", " ", t).strip()


def split_item(text: str) -> tuple[str, str]:
    """A short title for column B and the rest as notes for column D."""
    bold = re.match(r"\s*\*\*(.+?)\*\*[:.]?\s*(.*)", text, re.S)
    if bold:
        return clean(bold.group(1)).rstrip(":."), clean(bold.group(2))
    plain = clean(text)
    for sep in (" — ", ". ", ": "):
        i = plain.find(sep)
        if 0 < i <= 140:
            return plain[:i].rstrip("."), plain[i + len(sep):]
    return (plain[:140] + "…", plain[140:]) if len(plain) > 160 else (plain, "")


def parse():
    phases, phase = [], None
    for line in TODO.read_text(encoding="utf-8").splitlines():
        if line.startswith("## Phase"):
            phase = {"name": clean(line[3:]), "items": []}
            phases.append(phase)
            continue
        if line.startswith("## "):
            phase = None
            continue
        m = re.match(r"^\s*- \[( |x|X)\] (.*)$", line)
        if not (m and phase):
            continue
        done = m.group(1).lower() == "x"
        text = m.group(2)
        cancelled_phase = "CANCELLED" in phase["name"].upper()
        status = "Done" if done else ("Deferred" if cancelled_phase or DEFERRED.search(text) else "Remaining")
        title, notes = split_item(text)
        phase["items"].append({"title": title, "notes": notes[:1500], "status": status})
    return phases


def cell(ws, row, col, value, *, bold=False, fill=None, center=False, wrap=True, color=None, size=10, italic=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=FONT, size=size, bold=bold, italic=italic, color=color)
    if fill:
        c.fill = fill
    c.alignment = Alignment(horizontal="center" if center else None, vertical="top", wrap_text=wrap)
    c.border = BORDER
    return c


def header(ws, labels):
    for i, label in enumerate(labels, 1):
        cell(ws, 1, i, label, bold=True, fill=HEADER_FILL, center=True, color="FFFFFF", size=11)


def detail_sheet(ws, phases, only_open=False):
    header(ws, ["Phase", "Route / Item", "Status", "Notes" if not only_open else "Why it's not done / what's blocking it", "Phase name"])
    r = 2
    for ph in phases:
        items = [i for i in ph["items"] if not only_open or i["status"] != "Done"]
        if not items:
            continue
        for col in range(1, 5):
            cell(ws, r, col, ph["name"] if col == 1 else None, bold=True, fill=PHASE_FILL, size=11)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        r += 1
        for it in items:
            cell(ws, r, 1, None)
            cell(ws, r, 2, it["title"])
            cell(ws, r, 3, it["status"], fill=STATUS_FILL[it["status"]], center=True)
            cell(ws, r, 4, it["notes"])
            # Hidden helper column the Summary formulas count by.
            ws.cell(row=r, column=5, value=ph["name"]).font = Font(name=FONT, size=10)
            r += 1
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 100 if only_open else 95
    ws.column_dimensions["E"].hidden = True
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:D{r - 1}"
    return r - 1


def main():
    phases = parse()
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    full = wb.create_sheet("Full Detail")
    remaining = wb.create_sheet("Remaining Work")
    last_full = detail_sheet(full, phases)
    detail_sheet(remaining, phases, only_open=True)

    today = datetime.date.today().isoformat()
    cell(summary, 1, 1, "Tony Greenberg site — Next.js migration status", bold=True, size=16, wrap=False).border = Border()
    note = cell(
        summary, 2, 1,
        f"Generated from NEXTJS-MIGRATION-TODO.md on {today} by scripts/build-migration-status-report.py — "
        "a living document; re-run the script after updating the todo file. Counts are formulas over the Full Detail tab.",
        italic=True, color="666666", size=9, wrap=False,
    )
    note.border = Border()
    labels = ["Phase", "Done", "Remaining", "Deferred", "Total", "% Complete"]
    for i, label in enumerate(labels, 1):
        cell(summary, 4, i, label, bold=True, fill=HEADER_FILL, center=True, color="FFFFFF", size=11)

    rng = f"$E$2:$E${last_full}"
    st = f"$C$2:$C${last_full}"
    r = 5
    for ph in phases:
        cell(summary, r, 1, ph["name"])
        for col, status in ((2, "Done"), (3, "Remaining"), (4, "Deferred")):
            cell(summary, r, col, f"=COUNTIFS('Full Detail'!{rng},$A{r},'Full Detail'!{st},\"{status}\")", center=True)
        cell(summary, r, 5, f"=SUM(B{r}:D{r})", center=True)
        pct = cell(summary, r, 6, f"=IF(E{r}=0,0,B{r}/E{r})", center=True)
        pct.number_format = "0%"
        r += 1
    first, last = 5, r - 1
    cell(summary, r, 1, "TOTAL", bold=True, fill=PHASE_FILL)
    for col, letter in ((2, "B"), (3, "C"), (4, "D"), (5, "E")):
        cell(summary, r, col, f"=SUM({letter}{first}:{letter}{last})", bold=True, fill=PHASE_FILL, center=True)
    pct = cell(summary, r, 6, f"=IF(E{r}=0,0,B{r}/E{r})", bold=True, fill=PHASE_FILL, center=True)
    pct.number_format = "0%"
    summary.column_dimensions["A"].width = 52
    for letter in "BCDEF":
        summary.column_dimensions[letter].width = 12
    summary.freeze_panes = "A5"

    wb.save(OUT)
    done = sum(i["status"] == "Done" for p in phases for i in p["items"])
    total = sum(len(p["items"]) for p in phases)
    print(f"Wrote {OUT.name}: {len(phases)} phases, {total} items, {done} done")


if __name__ == "__main__":
    main()
